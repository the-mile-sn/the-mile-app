"""Read-only preview and deliberate bulk import of historical assessment spreadsheets."""
import base64
import io
import json
from datetime import date, datetime
from openpyxl import load_workbook
import assessments_v2

COLUMN_MAP={
 "E":"weight","N":"sum_8","O":"body_fat","AG":"bmi","AJ":"muscle_index",
 "F":"back","G":"triceps","H":"biceps","I":"suprailiac","J":"supraspinal",
 "K":"abdominal_skin","L":"thigh_skin","M":"calf_skin",
 "Q":"arm_right_relaxed","S":"chest","T":"waist","U":"abdominal_perimeter",
 "X":"thigh_right","AD":"thigh_left","Z":"calf_right","AF":"calf_left",
}
HEADERS={"A":"Data","B":"Nome","E":"Peso (kg)","N":"Σ8","O":"Massa gorda (Σ4)"}
LIMIT=5*1024*1024

def _decode(encoded):
 if not isinstance(encoded,str) or len(encoded)>LIMIT*2:raise ValueError("invalid_excel")
 try: raw=base64.b64decode(encoded,validate=True)
 except Exception:raise ValueError("invalid_excel")
 if len(raw)>LIMIT or not raw.startswith(b"PK"):raise ValueError("invalid_excel")
 return raw

def _date(value):
 if isinstance(value,datetime):return value.date()
 if isinstance(value,date):return value
 if isinstance(value,str):
  for fmt in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y"):
   try:return datetime.strptime(value.strip(),fmt).date()
   except ValueError:pass
 return None

def preview(encoded):
 raw=_decode(encoded)
 try: wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 except Exception:raise ValueError("invalid_excel")
 try:
  sheet=wb.worksheets[0]
  columns={col:__import__("openpyxl").utils.column_index_from_string(col)-1 for col in COLUMN_MAP}
  first=next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),())
  for col,label in HEADERS.items():
   i=__import__("openpyxl").utils.column_index_from_string(col)-1
   if i>=len(first) or str(first[i] or "").strip()!=label:raise ValueError("unsupported_excel_template")
  rows=[]
  seen=set()
  for number,rawrow in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
   if number>3000:raise ValueError("too_many_assessments")
   if not any(v is not None for v in rawrow):continue
   when=_date(rawrow[0] if len(rawrow)>0 else None)
   if not when:continue
   measurements={};warnings=[]
   for column,key in COLUMN_MAP.items():
    i=columns[column]
    value=rawrow[i] if i<len(rawrow) else None
    if value is None or str(value).strip()=="":continue
    try:
     if key=="body_fat" and isinstance(value,(int,float)) and 0<=value<=1:value=value*100
     normalized=assessments_v2.normalize({key:value})
     measurements.update(normalized)
    except ValueError:warnings.append("Valor inválido na coluna "+column)
   if when.isoformat() in seen:warnings.append("Data duplicada dentro do ficheiro")
   seen.add(when.isoformat())
   if when>date.today():warnings.append("Data futura: confirmar")
   dob=_date(rawrow[48] if len(rawrow)>48 else None)
   if dob and when<dob:warnings.append("Data anterior ao nascimento indicado: confirmar no Sheets")
   if not measurements:warnings.append("Sem medições compatíveis")
   rows.append({"row":number,"assessed_on":when.isoformat(),"name":str(rawrow[1] or "")[:100],
                "measurements":measurements,"warnings":warnings})
  if not rows:raise ValueError("no_assessments")
  return {"assessments":rows,"count":len(rows),"source_columns":len(first)}
 finally:wb.close()

def import_rows(c,athlete_id,records):
 if not isinstance(records,list) or not 0<len(records)<=3000:raise ValueError("invalid_assessments")
 existing={r["assessed_on"] for r in assessments_v2.rows(c,athlete_id)}
 saved=0;skipped=0
 for record in records:
  if not isinstance(record,dict):raise ValueError("invalid_assessments")
  day=record.get("assessed_on")
  try:date.fromisoformat(day)
  except (TypeError,ValueError):raise ValueError("invalid_date")
  if day in existing:skipped+=1;continue
  readings=assessments_v2.normalize(record.get("measurements",{}))
  if not readings:skipped+=1;continue
  assessments_v2.save(c,athlete_id,{"assessed_on":day,"measurements":readings,"averages":{},"formula_note":"Histórico importado de Excel","publish":True},None)
  saved+=1;existing.add(day)
 return {"imported":saved,"skipped":skipped}
